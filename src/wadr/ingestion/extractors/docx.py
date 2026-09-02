"""DOCX text extraction via python-docx."""

import io

from docx import Document


def extract(file_bytes: bytes) -> str:
    """Extract plain text from a .docx: paragraph text, then table cell text."""
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(p for p in parts if p.strip())
